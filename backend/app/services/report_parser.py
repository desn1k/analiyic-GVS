"""Парсер отчёта "Количество случаев нарушения качества ГВС (подробный)"."""
import re
from datetime import datetime, date
from typing import BinaryIO

import openpyxl

TITLE_RE = re.compile(r"за\s+(\d{2}\.\d{2}\.\d{2,4})\s*-\s*(\d{2}\.\d{2}\.\d{2,4})")
GENERATED_RE = re.compile(r"(\d{2}\.\d{2}\.\d{4})\s+(\d{2}:\d{2}:\d{2})")
CITY_RE = re.compile(r"(?:^|,)\s*([А-Яа-яЁё\-]+(?:\s[А-Яа-яЁё\-]+)*)\s+г(?:\s*[,(]|$)")

HEADER_ROWS = 4  # строки 1-4: заголовок отчёта, дата формирования, шапка таблицы (2 строки)

# Поля ищутся по подстроке в тексте заголовка (строка 3 — основная шапка,
# строка 4 — подзаголовки блока проверок). Колонки определяются динамически,
# а не по фиксированным номерам: разные выгрузки отчёта объединяют
# (merge) заголовочные ячейки по-разному (например, "ID ТУ" может занимать
# один или два столбца), из-за чего фиксированные индексы "плывут".
INT_FIELDS = {
    "total_records", "valid_records", "hours_total",
    "hours_violation_low", "hours_violation_high", "flow_below_2pct",
    "check_no_data", "check_vnr", "check_m1_m2", "check_qinj",
    "check_t_range", "check_t1_lt_t2", "check_v_max", "check_t1_const",
    "check_t_other",
}
FLOAT_FIELDS = {
    "avg_temp_ctp", "avg_temp_gvs", "contract_load", "volume_total",
    "volume_below_40", "volume_40_60", "volume_60_75", "volume_above_75",
}

# (имя поля, предикат(текст_заголовка) -> bool). Проверяются по порядку,
# первая подходящая ещё не занятая колонка побеждает.
FIELD_MATCHERS = [
    ("seq", lambda t: t == "№ пп"),
    ("object_name", lambda t: t == "Объект"),
    ("tu_name", lambda t: t == "ТУ"),
    ("object_id", lambda t: t == "ID объекта"),
    ("object_type", lambda t: t == "Тип объекта"),
    ("tu_id", lambda t: t == "ID ТУ"),
    ("source_name", lambda t: t == "ЦТП/Источник"),
    ("avg_temp_ctp", lambda t: "Средняя температура ГВС ЦТП" in t),
    ("is_dead_end", lambda t: "тупиковой системе" in t),
    ("system_type", lambda t: "закрытой системе" in t),
    ("total_records", lambda t: t == "Всего записей"),
    ("valid_records", lambda t: t == "Корректных записей"),
    ("avg_temp_gvs", lambda t: t.startswith("Средняя температура ГВС,")),
    ("contract_load", lambda t: "Договорная нагрузка" in t),
    ("volume_total", lambda t: t.startswith("Объём ГВС")),
    ("volume_below_40", lambda t: t.startswith("Т1 ГВС < 40")),
    ("volume_40_60", lambda t: "40 <= Т1 ГВС < 60" in t or "40<=Т1 ГВС<60" in t),
    ("volume_60_75", lambda t: "60 <= Т1 ГВС < 75" in t or "60<=Т1 ГВС<75" in t),
    ("volume_above_75", lambda t: t.startswith("Т1 ГВС >= 75") or t.startswith("Т1 ГВС>=75")),
    ("hours_total", lambda t: "часов периода" in t),
    ("hours_violation_low", lambda t: "нарушением" in t and "занижение" in t),
    ("hours_violation_high", lambda t: "нарушением" in t and "завышение" in t),
    ("flow_below_2pct", lambda t: "Расход менее" in t),
    ("check_no_data", lambda t: "н/д" in t),
    ("check_vnr", lambda t: "ВНР" in t),
    ("check_m1_m2", lambda t: "М1<0" in t or "M1<0" in t),
    ("check_qinj", lambda t: "Qинж" in t or "Qинж".lower() in t.lower()),
    ("check_t_range", lambda t: t.startswith("Т>=100") or t.startswith("T>=100")),
    ("check_t1_lt_t2", lambda t: "T1<T2" in t or "Т1<Т2" in t),
    ("check_v_max", lambda t: t.startswith("V>") or t.startswith("V >")),
    ("check_t1_const", lambda t: "const" in t),
    ("check_t_other", lambda t: t == "Проверка Т"),
    ("scheme", lambda t: t == "Схема"),
]


def extract_city(object_name: str) -> str | None:
    """Эвристика: ищет сегмент вида '<Город> г' в адресе объекта,
    пропуская сегменты вида '<Область> обл'."""
    if not object_name:
        return None
    for m in CITY_RE.finditer(object_name):
        name = m.group(1).strip()
        if "обл" not in name:
            return name
    first = object_name.split(",")[0].strip()
    return first or None


def _parse_date(s: str) -> date:
    s = s.strip()
    fmt = "%d.%m.%Y" if len(s.split(".")[-1]) == 4 else "%d.%m.%y"
    return datetime.strptime(s, fmt).date()


def _to_int(v):
    if v in (None, ""):
        return None
    try:
        return int(float(v))
    except (ValueError, TypeError):
        return None


def _to_float(v):
    if v in (None, ""):
        return None
    try:
        return float(v)
    except (ValueError, TypeError):
        return None


def _expand_merged_row(ws, row_idx: int, max_col: int) -> list[str]:
    """Возвращает текст заголовков строки row_idx, заполняя объединённые
    ячейки значением их левой верхней ячейки (по умолчанию в openpyxl все
    ячейки объединённого диапазона, кроме первой, имеют значение None)."""
    values = [ws.cell(row_idx, c).value for c in range(1, max_col + 1)]
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row_idx <= rng.max_row:
            anchor = ws.cell(rng.min_row, rng.min_col).value
            for c in range(rng.min_col, rng.max_col + 1):
                if rng.min_row <= row_idx <= rng.max_row:
                    values[c - 1] = anchor
    return [str(v).strip() if v is not None else "" for v in values]


def _detect_columns(ws) -> dict[str, int]:
    """Сопоставляет имена полей с индексами колонок (0-based) по тексту
    заголовков в строках 3 и 4, устойчиво к сдвигам из-за разного
    объединения ячеек между разными выгрузками отчёта."""
    max_col = ws.max_column
    top = _expand_merged_row(ws, 3, max_col)
    sub = _expand_merged_row(ws, 4, max_col)
    combined = [(sub[i] or top[i]) for i in range(max_col)]

    columns: dict[str, int] = {}
    used = set()
    for name, pred in FIELD_MATCHERS:
        for i, text in enumerate(combined):
            if i in used or not text:
                continue
            if pred(text):
                columns[name] = i
                used.add(i)
                break

    required = {"seq", "object_name", "tu_id", "total_records", "valid_records", "volume_total"}
    missing = required - columns.keys()
    if missing:
        raise ValueError(
            f"Не удалось определить колонки отчёта: {', '.join(sorted(missing))}. "
            "Похоже, формат файла отличается от ожидаемого."
        )
    return columns


def parse_report(file: BinaryIO):
    """Возвращает (period_start, period_end, generated_at, rows)."""
    wb_headers = openpyxl.load_workbook(file, data_only=True, read_only=False)
    ws_headers = wb_headers[wb_headers.sheetnames[0]]
    columns = _detect_columns(ws_headers)
    wb_headers.close()

    file.seek(0)
    wb = openpyxl.load_workbook(file, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]

    rows_iter = ws.iter_rows(values_only=True)
    title_row = next(rows_iter)
    title_cell = next((c for c in title_row if isinstance(c, str) and "за" in c), None)
    if not title_cell:
        raise ValueError("Не найден заголовок с периодом отчёта")
    m = TITLE_RE.search(title_cell)
    if not m:
        raise ValueError(f"Не удалось распознать период в заголовке: {title_cell!r}")
    period_start = _parse_date(m.group(1))
    period_end = _parse_date(m.group(2))

    gen_row = next(rows_iter)
    gen_cell = next((c for c in gen_row if isinstance(c, str) and "формирования" in c), None)
    generated_at = None
    if gen_cell:
        gm = GENERATED_RE.search(gen_cell)
        if gm:
            generated_at = datetime.strptime(f"{gm.group(1)} {gm.group(2)}", "%d.%m.%Y %H:%M:%S")

    next(rows_iter)  # строка заголовков (верхняя)
    next(rows_iter)  # строка заголовков (нижняя, подзаголовки проверок)

    rows = []
    for raw in rows_iter:
        if raw is None or all(v is None for v in raw):
            continue
        seq_val = raw[columns["seq"]]
        object_val = raw[columns["object_name"]]
        if seq_val is None or object_val is None:
            continue
        # Некоторые выгрузки склеивают несколько "печатных страниц" в один
        # лист, повторяя строку заголовков таблицы внутри данных — отсеиваем их.
        if isinstance(seq_val, str) and seq_val.strip() == "№ пп":
            continue
        if isinstance(object_val, str) and object_val.strip() == "Объект":
            continue
        row = {name: None for name, _ in FIELD_MATCHERS}
        for name, idx in columns.items():
            val = raw[idx] if idx < len(raw) else None
            if name in INT_FIELDS:
                val = _to_int(val)
            elif name in FLOAT_FIELDS:
                val = _to_float(val)
            elif isinstance(val, str):
                val = val.strip()
            row[name] = val
        rows.append(row)

    return period_start, period_end, generated_at, rows
