"""Импорт «Сводного отчёта по поиску конечного источника».

Каждая строка — точка учёта потребителя (E) с цепочкой вверх по уровням
(Уровень 1..7: ТК/ЦТП/теплоузлы) до конечного источника (котельная/ТЭЦ).
Строим цепочку узлов и связываем с потребителем по его tu_id (столбец E).
"""
import json

import openpyxl
from openpyxl.utils import column_index_from_string as ci
from sqlalchemy.orm import Session

from app.models import DeviceUpload, Hierarchy

COL_CONSUMER_TU = ci("E")     # ID точки учёта потребителя
COL_CONSUMER_OBJ = ci("I")    # ID объекта потребителя
LEVEL_COLS = [ci(c) for c in ("R", "AA", "AJ", "AS", "BB", "BK", "BT")]  # уровни 1..7
COL_SRC_TU = ci("CU")         # ID конечного источника
COL_SRC_ADDR = ci("CV")
COL_SRC_NAME = ci("CY")

# смещения внутри блока уровня (9 столбцов)
OFF_ADDR, OFF_NAME, OFF_TYPE = 1, 2, 4


def looks_like_hierarchy(file_obj) -> bool:
    try:
        file_obj.seek(0)
        wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        wb.close()
        if first and first[0] and "поиску конечного источника" in str(first[0]):
            return True
    except Exception:
        return False
    finally:
        file_obj.seek(0)
    return False


def _cell(row, idx):
    return row[idx - 1] if len(row) >= idx else None


def ingest_hierarchy(db: Session, source, upload_id: int) -> DeviceUpload:
    if hasattr(source, "seek"):
        source.seek(0)
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    upload = db.query(DeviceUpload).get(upload_id)

    seen: dict[str, dict] = {}
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i < 6:  # шапка на строках 1–5
            continue
        cons = _cell(row, COL_CONSUMER_TU)
        if not cons:
            continue
        cons = str(cons).strip()
        chain = []
        for lvl, col in enumerate(LEVEL_COLS, 1):
            tu = _cell(row, col)
            if not tu:
                continue
            chain.append({
                "level": lvl,
                "tu_id": str(tu).strip(),
                "address": _cell(row, col + OFF_ADDR),
                "name": _cell(row, col + OFF_NAME),
                "type": _cell(row, col + OFF_TYPE),
            })
        src = _cell(row, COL_SRC_TU)
        if src:
            chain.append({
                "level": "source",
                "tu_id": str(src).strip(),
                "address": _cell(row, COL_SRC_ADDR),
                "name": _cell(row, COL_SRC_NAME),
                "type": "Источник",
            })
        obj = _cell(row, COL_CONSUMER_OBJ)
        seen[cons] = {
            "consumer_tu_id": cons,
            "object_id": str(obj).strip() if obj else None,
            "chain": json.dumps(chain, ensure_ascii=False),
            "upload_id": upload_id,
        }
    wb.close()

    for cons, rec in seen.items():
        existing = db.query(Hierarchy).filter(Hierarchy.consumer_tu_id == cons).first()
        if existing is None:
            db.add(Hierarchy(**rec))
        else:
            existing.object_id = rec["object_id"]
            existing.chain = rec["chain"]
            existing.upload_id = upload_id

    upload.points_count = len(seen)
    upload.hours_count = len(seen)
    upload.status = "done"
    db.commit()
    return upload


def run_hierarchy_ingest_background(path: str, upload_id: int):
    import os
    from app.device_database import DeviceSessionLocal

    db = DeviceSessionLocal()
    try:
        ingest_hierarchy(db, path, upload_id)
    except Exception as e:  # noqa: BLE001
        db.rollback()
        up = db.query(DeviceUpload).get(upload_id)
        if up is not None:
            up.status = "error"
            up.error = str(e)[:500]
            db.commit()
    finally:
        db.close()
        try:
            os.remove(path)
        except OSError:
            pass
