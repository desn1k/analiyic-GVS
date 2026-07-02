"""Импорт «Ведомости отключений» (плановые/аварийные отключения ГВС).

Двухстрочная шапка (строки 1–2), данные с 3-й строки. Связь с объектами —
по нормализованному адресу (столбец N).
"""
from datetime import datetime
from typing import Optional

import openpyxl
from sqlalchemy.orm import Session

from app.models import DeviceUpload, Outage
from app.services.address import normalize_address

# 1-based индексы столбцов
C = {
    "number": 1, "kind": 2, "status": 3,
    "source": 11, "fias": 12, "guid": 13, "address": 14,
    "impact_plan": 15, "impact_fact": 16,
    "gvs": 18, "residents": 20, "reason": 24,
    "load_gvs": 27,
    "dur_plan": 29, "dur_fact": 30,
    "start_plan": 31, "start_fact": 32, "end_plan": 33, "end_fact": 34,
    "note": 37,
}

BATCH_SIZE = 2000


def looks_like_outage_report(file_obj) -> bool:
    try:
        file_obj.seek(0)
        wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
        if any("отключен" in s.lower() for s in wb.sheetnames):
            wb.close()
            return True
        ws = wb[wb.sheetnames[0]]
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        wb.close()
        if first and first[0] and "Номер отключения" in str(first[0]):
            return True
    except Exception:
        return False
    finally:
        file_obj.seek(0)
    return False


def _cell(row, key):
    idx = C[key] - 1
    return row[idx] if len(row) > idx else None


def _normalize_impact(v) -> Optional[str]:
    if not v:
        return None
    s = str(v).strip().lower()
    if "прекращ" in s:
        return "прекращение"
    if "ограничен" in s:
        return "ограничение"
    return "иное"


def _dt(v) -> Optional[datetime]:
    return v if isinstance(v, datetime) else None


def _num(v) -> Optional[float]:
    if isinstance(v, (int, float)):
        return float(v)
    return None


def _int(v) -> Optional[int]:
    if isinstance(v, (int, float)):
        return int(v)
    return None


def ingest_outage_report(db: Session, source, upload_id: int) -> DeviceUpload:
    if hasattr(source, "seek"):
        source.seek(0)
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    upload = db.query(DeviceUpload).get(upload_id)

    batch: list[dict] = []
    count = 0
    min_ts: Optional[datetime] = None
    max_ts: Optional[datetime] = None

    def flush(final=False):
        nonlocal batch
        if batch:
            db.bulk_insert_mappings(Outage, batch)
            batch = []
        if final or True:
            upload.hours_count = count
            db.commit()

    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i <= 2:  # двухстрочная шапка
            continue
        number = _cell(row, "number")
        address = _cell(row, "address")
        if number in (None, "") and not address:
            continue

        impact = _normalize_impact(_cell(row, "impact_fact") or _cell(row, "impact_plan"))
        start = _dt(_cell(row, "start_fact")) or _dt(_cell(row, "start_plan"))
        end = _dt(_cell(row, "end_fact")) or _dt(_cell(row, "end_plan"))

        rec = {
            "upload_id": upload_id,
            "number": str(number) if number is not None else None,
            "kind": _cell(row, "kind"),
            "status": _cell(row, "status"),
            "impact": impact,
            "address": str(address) if address else None,
            "address_norm": normalize_address(address),
            "fias": _cell(row, "fias"),
            "guid": _cell(row, "guid"),
            "source": _cell(row, "source"),
            "service_gvs": str(_cell(row, "gvs") or "").strip().lower() == "да",
            "reason": _cell(row, "reason"),
            "load_gkal": _num(_cell(row, "load_gvs")),
            "residents": _int(_cell(row, "residents")),
            "start_fact": _dt(_cell(row, "start_fact")),
            "end_fact": _dt(_cell(row, "end_fact")),
            "start_plan": _dt(_cell(row, "start_plan")),
            "end_plan": _dt(_cell(row, "end_plan")),
            "note": _cell(row, "note"),
        }
        batch.append(rec)
        count += 1
        if start and (min_ts is None or start < min_ts):
            min_ts = start
        if end and (max_ts is None or end > max_ts):
            max_ts = end
        if len(batch) >= BATCH_SIZE:
            flush()

    flush(final=True)
    wb.close()

    upload.period_start = min_ts
    upload.period_end = max_ts
    upload.points_count = count  # число строк ведомости
    upload.hours_count = count
    upload.status = "done"
    db.commit()
    return upload


def run_outage_ingest_background(path: str, upload_id: int):
    import os
    from app.device_database import DeviceSessionLocal

    db = DeviceSessionLocal()
    try:
        ingest_outage_report(db, path, upload_id)
    except Exception as e:  # noqa: BLE001
        db.rollback()
        up = db.query(DeviceUpload).get(upload_id)
        if up is not None:
            up.status = "error"
            up.error = str(e)[:500]
            db.commit()
    finally:
        db.close()
        try:
            os.remove(path)
        except OSError:
            pass
