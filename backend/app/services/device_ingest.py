"""Потоковый импорт почасовых данных с приборов учёта.

Формат «Ведомость учёта параметров потребления тепла в системе ГВС» —
стековые блоки по одному потребителю. В блоке: «Адрес:» (адрес), «Прибор
учёта:» (прибор), «Время на приборе:» (в столбце CV — UUID точки), затем
строка «Дата» и почасовые строки: A=дата, E=Т1, I=Т2, AI=V1, AV=V2 (объём).
Ниже — «Показания счётчиков» (нарастающие итоги), их пропускаем.

Число часов в периоде произвольное — читаем все строки с датой до следующего
маркера. Читаем строго потоково (openpyxl read_only) и пишем пачками, память
не зависит от размера файла.
"""
import re
from datetime import datetime
from typing import Optional

import openpyxl
from sqlalchemy import insert
from sqlalchemy.orm import Session

from app.models import DeviceUpload, DevicePoint, DeviceHourly

BATCH_SIZE = 5000
_TS_FORMATS = ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S")


def _num(v) -> Optional[float]:
    if v is None or v == "" or v == "---":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ".").strip())
    except ValueError:
        return None


def _parse_ts(v) -> Optional[datetime]:
    if isinstance(v, datetime):
        return v
    if not isinstance(v, str):
        return None
    s = v.strip()
    for fmt in _TS_FORMATS:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


# ---------------------------------------------------------------------------
# Второй приборный формат: «Ведомость учёта параметров потребления тепла в ГВС».
# Стековые блоки (~109 строк) по одному потребителю. UUID точки — в столбце CV
# на строке «Время на приборе:», адрес — «Адрес:» (B). Данные: A=дата, E=Т1,
# I=Т2, AI=V1 (объём, м³), AV=V2. «-» = недостоверно.
# ---------------------------------------------------------------------------
CONS_COLS = {"t1": 5, "t2": 9, "v1": 35, "v2": 48}  # E, I, AI, AV
CONS_UUID_COL = 100  # CV
CONS_ADDR_COL = 2    # B
CONS_DEV_COL = 4     # D
_CONS_TITLE_RE = re.compile(r"Ведомость\s+учёта\s+параметров\s+потреблени")


def looks_like_consumption_report(file_obj) -> bool:
    try:
        file_obj.seek(0)
        wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        wb.close()
        if first and first[0] and _CONS_TITLE_RE.search(str(first[0])):
            return True
    except Exception:
        return False
    finally:
        file_obj.seek(0)
    return False


def ingest_consumption_report(db: Session, source, upload_id: int) -> DeviceUpload:
    if hasattr(source, "seek"):
        source.seek(0)
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    upload = db.query(DeviceUpload).get(upload_id)

    cur_addr = cur_dev = cur_uuid = None
    in_data = False
    points: dict[str, dict] = {}
    batch: list[dict] = []
    hours = 0
    min_ts = max_ts = None
    cons_items = list(CONS_COLS.items())

    def flush(final=False):
        nonlocal batch
        if batch:
            db.execute(insert(DeviceHourly).prefix_with("OR REPLACE"), batch)
            batch = []
        if final or True:
            upload.hours_count = hours
            upload.points_count = len(points)
            db.commit()

    for row in ws.iter_rows(values_only=True):
        a = row[0]
        if isinstance(a, str):
            head = a.strip()
            if head.startswith("Адрес"):
                cur_addr = row[CONS_ADDR_COL - 1] if len(row) >= CONS_ADDR_COL else None
                in_data = False
                continue
            if head.startswith("Прибор учёта"):
                cur_dev = row[CONS_DEV_COL - 1] if len(row) >= CONS_DEV_COL else None
                in_data = False
                continue
            if head.startswith("Время на приборе"):
                ge = row[CONS_UUID_COL - 1] if len(row) >= CONS_UUID_COL else None
                cur_uuid = str(ge).strip() if ge not in (None, "") else None
                in_data = False
                continue
            if head == "Дата":
                in_data = True
                continue
            if head and not a[:1].isdigit():
                in_data = False  # «Показания счётчиков», «Ведомость учёта», и т.п.
                continue

        if not in_data or cur_uuid is None:
            continue
        ts = _parse_ts(a)
        if ts is None:
            continue

        rec = {"tu_uuid": cur_uuid, "upload_id": upload_id, "ts": ts}
        t1_raw = row[CONS_COLS["t1"] - 1] if len(row) >= CONS_COLS["t1"] else None
        for name, col in cons_items:
            raw = row[col - 1] if len(row) >= col else None
            rec[name] = _num(raw)
        # Достоверность ГВС считаем по температуре подачи t1; обратка t2 у ГВС
        # часто отсутствует («-») и не должна обнулять час.
        rec["valid"] = t1_raw not in (None, "", "-", "---") and rec.get("t1") is not None
        batch.append(rec)
        hours += 1
        if min_ts is None or ts < min_ts:
            min_ts = ts
        if max_ts is None or ts > max_ts:
            max_ts = ts
        if cur_uuid not in points:
            points[cur_uuid] = {
                "object_name": cur_addr, "device_name": cur_dev,
                "tu_name": cur_dev, "resource": "ГВС", "scheme": None,
            }
        if len(batch) >= BATCH_SIZE:
            flush()

    flush(final=True)
    wb.close()

    for uuid, meta in points.items():
        existing = db.query(DevicePoint).filter(DevicePoint.tu_uuid == uuid).first()
        if existing is None:
            db.add(DevicePoint(tu_uuid=uuid, last_upload_id=upload_id, **meta))
        else:
            for k, v in meta.items():
                setattr(existing, k, v)
            existing.last_upload_id = upload_id

    upload.period_start = min_ts
    upload.period_end = max_ts
    upload.points_count = len(points)
    upload.hours_count = hours
    upload.status = "done"
    db.commit()
    return upload


def run_device_ingest_background(path: str, upload_id: int):
    """Фоновый разбор: своя сессия БД, удаляет временный файл, ловит ошибки."""
    import os
    from app.device_database import DeviceSessionLocal

    db = DeviceSessionLocal()
    try:
        ingest_consumption_report(db, path, upload_id)
    except Exception as e:  # noqa: BLE001
        db.rollback()
        up = db.query(DeviceUpload).get(upload_id)
        if up is not None:
            up.status = "error"
            up.error = str(e)[:500]
            db.commit()
    finally:
        db.close()
        try:
            os.remove(path)
        except OSError:
            pass
