"""Импорт реестра объектов (паспорта из АИИС).

Одна строка на объект, ~102 столбца. Ключ — «Идентификатор объекта» (A),
совпадает с object_id недельного отчёта. Связываем паспорт с точками учёта
по object_id, дополнительно храним нормализованный адрес.
"""
from typing import Optional

import openpyxl
from openpyxl.utils import column_index_from_string as ci
from sqlalchemy.orm import Session

from app.models import DeviceUpload, ObjectRegistry
from app.services.address import normalize_address

C = {
    "object_id": "A", "address": "F", "t_supply": "J", "t_return": "K",
    "heat_system": "Z", "name": "AD", "q_heating": "AE", "q_gvs": "AF",
    "aiis_url": "CX", "fias": "CV",
}
IDX = {k: ci(v) - 1 for k, v in C.items()}
BATCH_SIZE = 2000


def looks_like_registry(file_obj) -> bool:
    try:
        file_obj.seek(0)
        wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        wb.close()
        if first and first[0] and "Идентификатор объекта" in str(first[0]):
            return True
    except Exception:
        return False
    finally:
        file_obj.seek(0)
    return False


def _cell(row, key):
    i = IDX[key]
    return row[i] if i < len(row) else None


def _num(v) -> Optional[float]:
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.replace(",", ".").strip())
        except ValueError:
            return None
    return None


def ingest_registry(db: Session, source, upload_id: int) -> DeviceUpload:
    if hasattr(source, "seek"):
        source.seek(0)
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    upload = db.query(DeviceUpload).get(upload_id)

    seen: dict[str, dict] = {}
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1:
            continue
        oid = _cell(row, "object_id")
        if not oid:
            continue
        oid = str(oid).strip()
        addr = _cell(row, "address")
        seen[oid] = {
            "object_id": oid,
            "upload_id": upload_id,
            "name": _cell(row, "name"),
            "address": str(addr) if addr else None,
            "address_norm": normalize_address(addr),
            "design_t_supply": _num(_cell(row, "t_supply")),
            "design_t_return": _num(_cell(row, "t_return")),
            "heat_system": _cell(row, "heat_system"),
            "q_heating": _num(_cell(row, "q_heating")),
            "q_gvs": _num(_cell(row, "q_gvs")),
            "aiis_url": _cell(row, "aiis_url"),
            "fias": _cell(row, "fias"),
        }
    wb.close()

    for oid, rec in seen.items():
        existing = db.query(ObjectRegistry).filter(ObjectRegistry.object_id == oid).first()
        if existing is None:
            db.add(ObjectRegistry(**rec))
        else:
            for k, v in rec.items():
                setattr(existing, k, v)

    upload.points_count = len(seen)
    upload.hours_count = len(seen)
    upload.status = "done"
    db.commit()
    return upload


def run_registry_ingest_background(path: str, upload_id: int):
    import os
    from app.device_database import DeviceSessionLocal

    db = DeviceSessionLocal()
    try:
        ingest_registry(db, path, upload_id)
    except Exception as e:  # noqa: BLE001
        db.rollback()
        up = db.query(DeviceUpload).get(upload_id)
        if up is not None:
            up.status = "error"
            up.error = str(e)[:500]
            db.commit()
    finally:
        db.close()
        try:
            os.remove(path)
        except OSError:
            pass
