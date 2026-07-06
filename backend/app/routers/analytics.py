from datetime import datetime
from typing import Optional

from fastapi import APIRouter, Depends, Query, HTTPException
from sqlalchemy.orm import Session

from app.database import get_db
from app.device_database import get_device_db
from app.models import (
    ReportPeriod, TuReportRow, Outage, ObjectRegistry, DevicePoint, DeviceHourly,
)
from app.schemas.schemas import (
    TuRowOut, DynamicsPoint, FilterOptions, WeeklySummary,
    ObjectComparisonOut, ObjectComparisonRow, ObjectPeriodMetric, PeriodOut,
)
from app.services.address import normalize_address, is_matchable_address

router = APIRouter(prefix="/api/analytics", tags=["analytics"])


def _apply_filters(q, object_type, is_dead_end, system_type, source_name, object_search):
    if object_type:
        q = q.filter(TuReportRow.object_type == object_type)
    if is_dead_end:
        q = q.filter(TuReportRow.is_dead_end == is_dead_end.strip().lower())
    if system_type:
        q = q.filter(TuReportRow.system_type == system_type)
    if source_name:
        q = q.filter(TuReportRow.source_name == source_name)
    return q


def _filter_by_search(rows, object_search):
    if not object_search:
        return rows
    needle = object_search.strip().lower()
    return [r for r in rows if r.object_name and needle in r.object_name.lower()]


def _violation_volume(row: TuReportRow) -> float:
    return (row.volume_below_40 or 0) + (row.volume_40_60 or 0) + (row.volume_above_75 or 0)


def _cap_pct(pct: float) -> float:
    return min(max(pct, 0.0), 100.0)


def _data_quality_pct(row: TuReportRow) -> float:
    pct = row.valid_records / row.total_records * 100 if row.total_records else 0.0
    return round(_cap_pct(pct), 1)


MIN_MEANINGFUL_VOLUME = 0.05  # м³ — ниже этого порога объём считается шумом измерения, не нарушением


def _probable_cause(
    below_40: float, v40_60: float, above_75: float,
    hours_violation_low: int, hours_violation_high: int, hours_total: int,
    is_dead_end: Optional[str], data_quality_pct: float,
    volume_total: float = None,
) -> str:
    """Эвристика (rule-based) для определения вероятной причины нарушения по объёмам/часам.
    При появлении данных приборов (почасовые ряды) можно заменить/дополнить ML-моделью."""
    violation = below_40 + v40_60 + above_75
    if violation <= 0:
        return "Норма"
    if volume_total is not None and volume_total < MIN_MEANINGFUL_VOLUME:
        return "Недостаточно объёма для оценки"
    if data_quality_pct < 52:
        return "Недостаточно данных для определения причины"

    buckets = {"below_40": below_40, "40_60": v40_60, "above_75": above_75}
    dominant = max(buckets, key=buckets.get)

    if dominant == "above_75":
        cause = "Перегрев — завышенная температура подачи"
    elif dominant == "below_40":
        cause = "Сильный недогрев — ниже норматива"
    else:
        cause = "Погранично занижение — близко к нижней границе нормы"

    if dominant != "above_75" and str(is_dead_end or "").strip().lower() == "да":
        cause += " (тупиковая ветка — возможен дефицит циркуляции)"

    violation_hours = hours_violation_low + hours_violation_high
    if hours_total and violation_hours / hours_total > 0.5:
        cause += f"; нарушение хроническое ({violation_hours} ч. из {hours_total} ч. периода)"

    return cause


def _to_out(row: TuReportRow) -> TuRowOut:
    out = TuRowOut.model_validate(row)
    violation = _violation_volume(row)
    out.violation_volume = round(violation, 3)
    out.violation_pct = (
        round(_cap_pct(violation / row.volume_total * 100), 2)
        if row.volume_total and row.volume_total >= MIN_MEANINGFUL_VOLUME else 0.0
    )
    out.data_quality_pct = _data_quality_pct(row)
    out.probable_cause = _probable_cause(
        row.volume_below_40 or 0, row.volume_40_60 or 0, row.volume_above_75 or 0,
        row.hours_violation_low or 0, row.hours_violation_high or 0, row.hours_total or 0,
        row.is_dead_end, out.data_quality_pct, row.volume_total or 0,
    )
    return out


def _outage_index(device_db: Session):
    """Индексы отключений ГВС: по ФИАС и по нормализованному адресу."""
    fias_map: dict[str, set] = {}
    addr_map: dict[str, set] = {}
    q = device_db.query(Outage.fias, Outage.address_norm, Outage.impact).filter(
        Outage.service_gvs == True  # noqa: E712
    )
    for fias, addr, impact in q:
        imp = impact or "иное"
        if fias:
            fias_map.setdefault(fias, set()).add(imp)
        if addr:
            addr_map.setdefault(addr, set()).add(imp)
    return fias_map, addr_map


def _outage_maps(device_db: Session):
    """Готовит связь object_id -> типы отключений через реестр (ФИАС/адрес),
    а также прямую карту по адресу (запасной вариант без реестра)."""
    fias_map, addr_map = _outage_index(device_db)
    by_object: dict[str, set] = {}
    for reg in device_db.query(ObjectRegistry).all():
        impacts = set()
        if reg.fias and reg.fias in fias_map:
            impacts |= fias_map[reg.fias]
        if reg.address_norm and reg.address_norm in addr_map:
            impacts |= addr_map[reg.address_norm]
        if impacts:
            by_object[reg.object_id] = impacts
    return by_object, addr_map


def _tag_outages(rows: list[TuRowOut], maps) -> None:
    by_object, addr_map = maps
    for r in rows:
        impacts = by_object.get(r.object_id)
        if not impacts:
            addr = normalize_address(r.object_name)
            if is_matchable_address(addr):
                impacts = addr_map.get(addr)
        if impacts:
            r.has_outage = True
            r.outage_impacts = ", ".join(sorted(impacts))


def _filter_by_outage(rows: list[TuRowOut], outage: str) -> list[TuRowOut]:
    if outage == "any":
        return [r for r in rows if r.has_outage]
    return [r for r in rows if outage in (r.outage_impacts or "").split(", ")]


def _registry_map(device_db: Session) -> dict:
    """Карта object_id -> паспорт объекта из реестра."""
    return {r.object_id: r for r in device_db.query(ObjectRegistry).all()}


def _tag_registry(rows: list[TuRowOut], rmap: dict) -> None:
    for r in rows:
        reg = rmap.get(r.object_id)
        if not reg:
            continue
        r.registry_name = reg.name
        r.heat_system = reg.heat_system
        r.design_t_supply = reg.design_t_supply
        r.design_t_return = reg.design_t_return
        r.q_heating = reg.q_heating
        r.q_gvs = reg.q_gvs
        r.aiis_url = reg.aiis_url


@router.get("/periods/{period_id}/tu", response_model=list[TuRowOut])
def list_tu_rows(
    period_id: int,
    object_type: Optional[str] = None,
    is_dead_end: Optional[str] = None,
    system_type: Optional[str] = None,
    source_name: Optional[str] = None,
    search: Optional[str] = None,
    min_violation_pct: Optional[float] = None,
    min_data_quality_pct: Optional[float] = None,
    max_data_quality_pct: Optional[float] = None,
    outage: Optional[str] = Query(None, pattern="^(any|прекращение|ограничение|иное)$"),
    sort_by: str = Query("violation_pct", pattern="^(violation_pct|volume_total|avg_temp_gvs|data_quality_pct)$"),
    order: str = Query("desc", pattern="^(asc|desc)$"),
    limit: int = 2000,
    db: Session = Depends(get_db),
    device_db: Session = Depends(get_device_db),
):
    q = db.query(TuReportRow).filter(TuReportRow.period_id == period_id)
    q = _apply_filters(q, object_type, is_dead_end, system_type, source_name, None)
    db_rows = _filter_by_search(q.all(), search)
    rows = [_to_out(r) for r in db_rows]
    _tag_outages(rows, _outage_maps(device_db))
    _tag_registry(rows, _registry_map(device_db))

    if min_violation_pct is not None:
        rows = [r for r in rows if r.violation_pct >= min_violation_pct]
    if min_data_quality_pct is not None:
        rows = [r for r in rows if r.data_quality_pct >= min_data_quality_pct]
    if max_data_quality_pct is not None:
        rows = [r for r in rows if r.data_quality_pct <= max_data_quality_pct]
    if outage:
        rows = _filter_by_outage(rows, outage)

    rows.sort(key=lambda r: getattr(r, sort_by) or 0, reverse=(order == "desc"))
    return rows[:limit]


@router.get("/tu", response_model=list[TuRowOut])
def list_all_tu_rows(
    object_type: Optional[str] = None,
    is_dead_end: Optional[str] = None,
    system_type: Optional[str] = None,
    source_name: Optional[str] = None,
    search: Optional[str] = None,
    min_violation_pct: Optional[float] = None,
    min_data_quality_pct: Optional[float] = None,
    max_data_quality_pct: Optional[float] = None,
    outage: Optional[str] = Query(None, pattern="^(any|прекращение|ограничение|иное)$"),
    sort_by: str = Query("violation_pct", pattern="^(violation_pct|volume_total|avg_temp_gvs|data_quality_pct)$"),
    order: str = Query("desc", pattern="^(asc|desc)$"),
    limit: int = 5000,
    db: Session = Depends(get_db),
    device_db: Session = Depends(get_device_db),
):
    """Все точки учёта из базы (без привязки к периоду).

    Одна строка на ТУ — берётся запись из самого свежего периода.
    """
    q = db.query(TuReportRow).join(ReportPeriod).order_by(ReportPeriod.period_start)
    q = _apply_filters(q, object_type, is_dead_end, system_type, source_name, None)
    db_rows = _filter_by_search(q.all(), search)

    # Дедуп по ТУ: более поздний период перезаписывает ранний.
    by_tu: dict = {}
    for r in db_rows:
        by_tu[r.tu_id or f"id{r.id}"] = r
    db_rows = list(by_tu.values())

    rows = [_to_out(r) for r in db_rows]
    _tag_outages(rows, _outage_maps(device_db))
    _tag_registry(rows, _registry_map(device_db))

    if min_violation_pct is not None:
        rows = [r for r in rows if r.violation_pct >= min_violation_pct]
    if min_data_quality_pct is not None:
        rows = [r for r in rows if r.data_quality_pct >= min_data_quality_pct]
    if max_data_quality_pct is not None:
        rows = [r for r in rows if r.data_quality_pct <= max_data_quality_pct]
    if outage:
        rows = _filter_by_outage(rows, outage)

    rows.sort(key=lambda r: getattr(r, sort_by) or 0, reverse=(order == "desc"))
    return rows[:limit]


@router.get("/periods/{period_id}/filters", response_model=FilterOptions)
def get_filters(period_id: int, db: Session = Depends(get_db)):
    q = db.query(TuReportRow).filter(TuReportRow.period_id == period_id)
    object_types = sorted({r[0] for r in q.with_entities(TuReportRow.object_type) if r[0]})
    system_types = sorted({r[0] for r in q.with_entities(TuReportRow.system_type) if r[0]})
    sources = sorted({r[0] for r in q.with_entities(TuReportRow.source_name) if r[0]})
    return FilterOptions(
        object_types=object_types, system_types=system_types, sources=sources,
    )


@router.get("/dynamics", response_model=list[DynamicsPoint])
def get_dynamics(
    object_type: Optional[str] = None,
    is_dead_end: Optional[str] = None,
    system_type: Optional[str] = None,
    source_name: Optional[str] = None,
    object_id: Optional[str] = None,
    db: Session = Depends(get_db),
):
    periods = db.query(ReportPeriod).order_by(ReportPeriod.period_start).all()
    points = []
    for p in periods:
        q = db.query(TuReportRow).filter(TuReportRow.period_id == p.id)
        q = _apply_filters(q, object_type, is_dead_end, system_type, source_name, None)
        if object_id:
            q = q.filter(TuReportRow.object_id == object_id)
        rows = q.all()
        if not rows:
            continue

        volume_total = sum(r.volume_total or 0 for r in rows)
        violation_volume = sum(_violation_volume(r) for r in rows)
        temp_weighted = sum((r.avg_temp_gvs or 0) * (r.volume_total or 0) for r in rows)
        avg_temp = temp_weighted / volume_total if volume_total else 0.0

        points.append(DynamicsPoint(
            period_id=p.id,
            period_start=p.period_start,
            period_end=p.period_end,
            volume_total=round(volume_total, 2),
            violation_volume=round(violation_volume, 2),
            violation_pct=round(_cap_pct(violation_volume / volume_total * 100), 2) if volume_total else 0.0,
            avg_temp_gvs=round(avg_temp, 2),
            hours_violation_low=sum(r.hours_violation_low or 0 for r in rows),
            hours_violation_high=sum(r.hours_violation_high or 0 for r in rows),
            tu_count=len(rows),
        ))
    return points


@router.get("/weekly-summary", response_model=list[WeeklySummary])
def get_weekly_summary(
    object_type: Optional[str] = None,
    is_dead_end: Optional[str] = None,
    system_type: Optional[str] = None,
    source_name: Optional[str] = None,
    db: Session = Depends(get_db),
):
    """Сводка по неделям: объекты агрегируются (несколько ТУ на объект суммируются)."""
    periods = db.query(ReportPeriod).order_by(ReportPeriod.period_start).all()
    result = []
    for p in periods:
        q = db.query(TuReportRow).filter(TuReportRow.period_id == p.id)
        q = _apply_filters(q, object_type, is_dead_end, system_type, source_name, None)
        rows = q.all()
        if not rows:
            continue

        by_object: dict[str, dict] = {}
        for r in rows:
            agg = by_object.setdefault(r.object_id, {"volume_total": 0.0, "violation_volume": 0.0, "above_75": 0.0})
            agg["volume_total"] += r.volume_total or 0
            agg["violation_volume"] += _violation_volume(r)
            agg["above_75"] += r.volume_above_75 or 0

        volume_total = sum(o["volume_total"] for o in by_object.values())
        violation_volume = sum(o["violation_volume"] for o in by_object.values())
        objects_with_violation = sum(1 for o in by_object.values() if o["violation_volume"] > 0)
        overheat_count = sum(1 for o in by_object.values() if o["above_75"] > 0)

        result.append(WeeklySummary(
            period_id=p.id,
            period_start=p.period_start,
            period_end=p.period_end,
            objects_count=len(by_object),
            tu_count=len(rows),
            objects_with_violation=objects_with_violation,
            volume_total=round(volume_total, 2),
            violation_volume=round(violation_volume, 2),
            violation_pct=round(_cap_pct(violation_volume / volume_total * 100), 2) if volume_total else 0.0,
            overheat_count=overheat_count,
        ))
    return result


@router.get("/object-comparison", response_model=ObjectComparisonOut)
def get_object_comparison(
    object_type: Optional[str] = None,
    is_dead_end: Optional[str] = None,
    system_type: Optional[str] = None,
    source_name: Optional[str] = None,
    search: Optional[str] = None,
    min_data_quality_pct: Optional[float] = None,
    max_data_quality_pct: Optional[float] = None,
    db: Session = Depends(get_db),
):
    """Сравнение объектов по неделям: одна строка на объект, столбцы — периоды."""
    periods = db.query(ReportPeriod).order_by(ReportPeriod.period_start).all()
    objects: dict[str, ObjectComparisonRow] = {}

    for p in periods:
        q = db.query(TuReportRow).filter(TuReportRow.period_id == p.id)
        q = _apply_filters(q, object_type, is_dead_end, system_type, source_name, None)
        rows = _filter_by_search(q.all(), search)
        if min_data_quality_pct is not None:
            rows = [r for r in rows if _data_quality_pct(r) >= min_data_quality_pct]
        if max_data_quality_pct is not None:
            rows = [r for r in rows if _data_quality_pct(r) <= max_data_quality_pct]

        by_object: dict[str, dict] = {}
        for r in rows:
            agg = by_object.setdefault(r.object_id, {
                "object_name": r.object_name, "object_type": r.object_type,
                "volume_total": 0.0, "violation_volume": 0.0,
                "below_40": 0.0, "v40_60": 0.0, "above_75": 0.0,
                "temp_weighted": 0.0, "total_records": 0, "valid_records": 0,
                "hours_total": 0, "hours_violation_low": 0, "hours_violation_high": 0,
                "is_dead_end": r.is_dead_end,
            })
            agg["volume_total"] += r.volume_total or 0
            agg["violation_volume"] += _violation_volume(r)
            agg["below_40"] += r.volume_below_40 or 0
            agg["v40_60"] += r.volume_40_60 or 0
            agg["above_75"] += r.volume_above_75 or 0
            agg["temp_weighted"] += (r.avg_temp_gvs or 0) * (r.volume_total or 0)
            agg["total_records"] += r.total_records or 0
            agg["valid_records"] += r.valid_records or 0
            agg["hours_total"] += r.hours_total or 0
            agg["hours_violation_low"] += r.hours_violation_low or 0
            agg["hours_violation_high"] += r.hours_violation_high or 0

        for object_id, agg in by_object.items():
            row = objects.get(object_id)
            if row is None:
                row = ObjectComparisonRow(
                    object_id=object_id,
                    object_name=agg["object_name"],
                    object_type=agg["object_type"],
                    periods={},
                )
                objects[object_id] = row

            vt = agg["volume_total"]
            data_quality_pct = round(_cap_pct(agg["valid_records"] / agg["total_records"] * 100), 1) if agg["total_records"] else 0.0
            row.periods[p.id] = ObjectPeriodMetric(
                period_id=p.id,
                volume_total=round(vt, 2),
                violation_volume=round(agg["violation_volume"], 2),
                violation_pct=(
                    round(_cap_pct(agg["violation_volume"] / vt * 100), 2)
                    if vt >= MIN_MEANINGFUL_VOLUME else 0.0
                ),
                avg_temp_gvs=round(agg["temp_weighted"] / vt, 2) if vt else 0.0,
                data_quality_pct=data_quality_pct,
                has_overheat=agg["above_75"] > 0,
                probable_cause=_probable_cause(
                    agg["below_40"], agg["v40_60"], agg["above_75"],
                    agg["hours_violation_low"], agg["hours_violation_high"], agg["hours_total"],
                    agg["is_dead_end"], data_quality_pct, vt,
                ),
            )

    period_counts = {}
    for p in periods:
        tu_count = (
            db.query(TuReportRow)
            .filter(TuReportRow.period_id == p.id)
            .count()
        )
        period_counts[p.id] = tu_count

    return ObjectComparisonOut(
        periods=[
            PeriodOut(
                id=p.id, period_start=p.period_start, period_end=p.period_end,
                generated_at=p.generated_at, source_filename=p.source_filename,
                tu_count=period_counts[p.id],
            )
            for p in periods
        ],
        objects=list(objects.values()),
    )


# ---------------------------------------------------------------------------
# Анализ качества ГВС по приборным почасовым данным (СанПиН с ночной поправкой).
# Период и пороги задаются вручную. Ночь 00:00–05:00: допускается снижение до
# −5°C (порог 55), днём −3°C (порог 57), перегрев >75. Исключаем недостоверные
# часы и часы, попавшие в интервалы отключений ГВС.
# ---------------------------------------------------------------------------
def _outage_intervals(device_db: Session) -> dict[str, list]:
    intervals: dict[str, list] = {}
    q = device_db.query(Outage.address_norm, Outage.start_fact, Outage.end_fact).filter(
        Outage.service_gvs == True  # noqa: E712
    )
    for addr, start, end in q:
        if addr and start and end:
            intervals.setdefault(addr, []).append((start, end))
    return intervals


def _low_threshold(ts, night_start, night_end, night_low, day_low):
    return night_low if night_start <= ts.hour < night_end else day_low


def _compute_gvs_quality(
    db, device_db, dt_from, dt_to, chronic_pct, min_reliability_pct,
    night_start, night_end, day_low, night_low, high,
    exclude_no_draw, scope, only_violations, limit,
):
    points = {p.tu_uuid: p for p in device_db.query(DevicePoint).all()}
    outage_map = _outage_intervals(device_db)

    # Точки учёта из базы ГВС (недельный отчёт): анализируем именно их, а не все
    # приборные точки (среди которых ЦТП/источники). Имя и признак тупика — из отчёта.
    report_info = {}
    for tu_id, oname, oid, is_de in db.query(
        TuReportRow.tu_id, TuReportRow.object_name, TuReportRow.object_id, TuReportRow.is_dead_end
    ).distinct():
        if tu_id and tu_id not in report_info:
            report_info[tu_id] = {"object_name": oname, "object_id": oid, "is_dead_end": is_de}
    scope_ids = set(report_info) if scope == "report" else None

    # цепочки иерархии (для атрибуции причины)
    from app.models import Hierarchy  # локальный импорт, чтобы не тянуть в топ
    import json
    hier_chain = {
        h.consumer_tu_id: json.loads(h.chain or "[]")
        for h in device_db.query(Hierarchy).all()
    }

    rows = (
        device_db.query(
            DeviceHourly.tu_uuid, DeviceHourly.ts, DeviceHourly.t1,
            DeviceHourly.v1, DeviceHourly.m1, DeviceHourly.valid
        )
        .filter(DeviceHourly.ts >= dt_from, DeviceHourly.ts <= dt_to)
        .all()
    )

    # почасовой t1 по всем точкам (для проверки вышестоящих ТУ)
    t1_by_tu: dict[str, dict] = {}
    agg: dict[str, dict] = {}
    for tu, ts, t1, v1, m1, valid in rows:
        if valid and t1 is not None:
            t1_by_tu.setdefault(tu, {})[ts] = t1
        if scope_ids is not None and tu not in scope_ids:
            continue
        a = agg.get(tu)
        if a is None:
            addr = normalize_address(points[tu].object_name) if tu in points else ""
            a = agg[tu] = {
                "total": 0, "valid": 0, "counted": 0, "excluded_outage": 0,
                "excluded_noflow": 0, "viol_low": 0, "viol_high": 0, "t1_sum": 0.0,
                "viol_ts": [], "intervals": outage_map.get(addr, []),
            }
        a["total"] += 1
        if not valid or t1 is None:
            continue
        a["valid"] += 1
        if any(s <= ts <= e for s, e in a["intervals"]):
            a["excluded_outage"] += 1
            continue
        vol = v1 if v1 is not None else m1
        if exclude_no_draw and (vol is None or vol <= 0.001):
            a["excluded_noflow"] += 1
            continue
        a["counted"] += 1
        a["t1_sum"] += t1
        low = _low_threshold(ts, night_start, night_end, night_low, day_low)
        if t1 < low:
            a["viol_low"] += 1
            a["viol_ts"].append(ts)
        elif t1 > high:
            a["viol_high"] += 1

    def _attribute(tu, viol_ts):
        """Системная (источник/сеть) или локальная причина недогрева — по иерархии."""
        chain = hier_chain.get(tu) or []
        upstream = None
        for n in chain:
            uid = n.get("tu_id")
            if uid and uid != tu and uid in t1_by_tu:
                upstream = (uid, n)
                break
        if not upstream or not viol_ts:
            return "источник выше по иерархии без данных — причина не определена"
        uid, node = upstream
        umap = t1_by_tu[uid]
        checked = same_low = 0
        for ts in viol_ts:
            ut = umap.get(ts)
            if ut is None:
                continue
            checked += 1
            low = _low_threshold(ts, night_start, night_end, night_low, day_low)
            if ut < low:
                same_low += 1
        if checked == 0:
            return "нет совпадающих часов у вышестоящей ТУ"
        label = node.get("name") or node.get("address") or uid
        if same_low / checked >= 0.5:
            return f"системная: недогрев и выше по сети ({label})"
        return f"локальная: выше по сети норма ({label}) — внутридомовая/ветка"

    result = []
    for tu, a in agg.items():
        total = a["total"]
        reliability = round(a["valid"] / total * 100, 1) if total else 0.0
        counted = a["counted"]
        viol = a["viol_low"] + a["viol_high"]
        viol_pct = round(viol / counted * 100, 1) if counted else 0.0
        info = report_info.get(tu, {})
        name = info.get("object_name") or (points[tu].object_name if tu in points else None)

        cause = ""
        if reliability < min_reliability_pct:
            verdict = "Недостаточно достоверных данных"
        elif counted == 0:
            verdict = "Нет зачтённых часов (всё исключено)"
        elif viol == 0:
            verdict = "Норма"
        elif a["viol_high"] > a["viol_low"]:
            verdict = "Перегрев (>75°C)"
        else:
            verdict = "Недогрев (ниже норматива)"
            cause = _attribute(tu, a["viol_ts"])
            if str(info.get("is_dead_end") or "").strip().lower() == "да":
                cause += "; тупиковая ветка"
        if viol and viol_pct > chronic_pct:
            verdict += f"; хроническое ({viol}/{counted} ч)"

        if only_violations and verdict in (
            "Норма", "Недостаточно достоверных данных", "Нет зачтённых часов (всё исключено)"
        ):
            continue

        result.append({
            "tu_uuid": tu,
            "object_name": name,
            "object_id": info.get("object_id"),
            "hours_total": total,
            "hours_valid": a["valid"],
            "hours_counted": counted,
            "hours_excluded_outage": a["excluded_outage"],
            "hours_excluded_noflow": a["excluded_noflow"],
            "reliability_pct": reliability,
            "viol_low": a["viol_low"],
            "viol_high": a["viol_high"],
            "viol_pct": viol_pct,
            "avg_t1": round(a["t1_sum"] / counted, 1) if counted else None,
            "verdict": verdict,
            "cause": cause,
        })

    result.sort(key=lambda x: x["viol_pct"], reverse=True)
    return {
        "date_from": dt_from, "date_to": dt_to,
        "params": {
            "chronic_pct": chronic_pct, "min_reliability_pct": min_reliability_pct,
            "night": [night_start, night_end], "day_low": day_low,
            "night_low": night_low, "high": high,
            "exclude_no_draw": exclude_no_draw, "scope": scope,
        },
        "count": len(result),
        "rows": result[:limit],
    }


# Общие query-параметры анализа качества ГВС.
def _quality_params(
    date_from: str = Query(..., description="ISO datetime начала периода"),
    date_to: str = Query(..., description="ISO datetime конца периода"),
    chronic_pct: float = 50.0,
    min_reliability_pct: float = 52.0,
    night_start: int = 0,
    night_end: int = 5,
    day_low: float = 57.0,
    night_low: float = 55.0,
    high: float = 75.0,
    exclude_no_draw: bool = True,
    scope: str = Query("report", pattern="^(report|all)$"),
    only_violations: bool = False,
    limit: int = 100000,
):
    try:
        dt_from = datetime.fromisoformat(date_from)
        dt_to = datetime.fromisoformat(date_to)
    except ValueError:
        raise HTTPException(400, "Неверный формат даты (ожидается ISO)")
    return dict(
        dt_from=dt_from, dt_to=dt_to, chronic_pct=chronic_pct,
        min_reliability_pct=min_reliability_pct, night_start=night_start,
        night_end=night_end, day_low=day_low, night_low=night_low, high=high,
        exclude_no_draw=exclude_no_draw, scope=scope,
        only_violations=only_violations, limit=limit,
    )


@router.get("/gvs-quality")
def gvs_quality(
    params: dict = Depends(_quality_params),
    db: Session = Depends(get_db),
    device_db: Session = Depends(get_device_db),
):
    return _compute_gvs_quality(db, device_db, **params)


# Заголовки столбцов отчёта (в порядке вывода).
_EXPORT_COLUMNS = [
    ("object_name", "Объект"),
    ("object_id", "ID объекта"),
    ("tu_uuid", "ID точки учёта"),
    ("hours_total", "Часов всего"),
    ("hours_valid", "Достоверных часов"),
    ("hours_counted", "Зачтено часов"),
    ("hours_excluded_outage", "Исключено (отключения)"),
    ("hours_excluded_noflow", "Исключено (без водоразбора)"),
    ("reliability_pct", "Достоверность, %"),
    ("viol_low", "Недогрев, ч"),
    ("viol_high", "Перегрев, ч"),
    ("viol_pct", "% нарушения"),
    ("avg_t1", "Сред. T подачи, °C"),
    ("verdict", "Вердикт"),
    ("cause", "Причина"),
]


@router.get("/gvs-quality/export")
def gvs_quality_export(
    params: dict = Depends(_quality_params),
    db: Session = Depends(get_db),
    device_db: Session = Depends(get_device_db),
):
    import io
    import openpyxl
    from openpyxl.styles import Font
    from fastapi.responses import StreamingResponse

    data = _compute_gvs_quality(db, device_db, **params)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Качество ГВС"

    df = data["date_from"].strftime("%d.%m.%Y %H:%M")
    dt = data["date_to"].strftime("%d.%m.%Y %H:%M")
    ws.append([f"Анализ качества ГВС за период {df} — {dt}"])
    ws["A1"].font = Font(bold=True, size=12)
    p = data["params"]
    ws.append([
        f"Пороги: день ≥{p['day_low']}°C, ночь {p['night'][0]}–{p['night'][1]} ≥{p['night_low']}°C, "
        f"перегрев >{p['high']}°C; хроническое >{p['chronic_pct']}%; достоверность ≥{p['min_reliability_pct']}%; "
        f"охват: {'база ГВС' if p['scope'] == 'report' else 'все приборные точки'}."
    ])
    ws.append([])

    header_row = ws.max_row + 1
    ws.append([title for _, title in _EXPORT_COLUMNS])
    for cell in ws[header_row]:
        cell.font = Font(bold=True)

    for r in data["rows"]:
        ws.append([r.get(key) for key, _ in _EXPORT_COLUMNS])

    # ширины столбцов
    for i, (_, title) in enumerate(_EXPORT_COLUMNS, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(12, min(46, len(title) + 4))
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"gvs_quality_{data['date_from'].strftime('%Y%m%d')}_{data['date_to'].strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


# ---------------------------------------------------------------------------
# Разбивка недели по источникам (ЦТП/котельные) и динамика одного источника.
# ---------------------------------------------------------------------------
def _final_source_map(device_db: Session) -> dict:
    """Карта tu_id потребителя -> название конечного источника (ТЭЦ/котельная) из иерархии."""
    import json
    from app.models import Hierarchy
    result = {}
    for h in device_db.query(Hierarchy).all():
        chain = json.loads(h.chain or "[]")
        for n in reversed(chain):
            if n.get("level") == "source":
                result[h.consumer_tu_id] = n.get("name") or n.get("address")
                break
    return result


@router.get("/weekly-summary/{period_id}/by-source")
def weekly_summary_by_source(
    period_id: int,
    object_type: Optional[str] = None,
    is_dead_end: Optional[str] = None,
    system_type: Optional[str] = None,
    db: Session = Depends(get_db),
    device_db: Session = Depends(get_device_db),
):
    q = db.query(TuReportRow).filter(TuReportRow.period_id == period_id)
    q = _apply_filters(q, object_type, is_dead_end, system_type, None, None)
    rows = q.all()
    final_map = _final_source_map(device_db)

    by_source: dict = {}
    for r in rows:
        src = r.source_name or "— без источника"
        s = by_source.setdefault(src, {"volume": 0.0, "viol": 0.0, "tu": 0, "objects": {}, "finals": {}})
        s["volume"] += r.volume_total or 0
        s["viol"] += _violation_volume(r)
        s["tu"] += 1
        o = s["objects"].setdefault(r.object_id, {"vv": 0.0, "a75": 0.0})
        o["vv"] += _violation_volume(r)
        o["a75"] += r.volume_above_75 or 0
        fin = final_map.get(r.tu_id)
        if fin:
            s["finals"][fin] = s["finals"].get(fin, 0) + 1

    result = []
    finals_set = set()
    for src, s in by_source.items():
        objs = s["objects"]
        final_source = max(s["finals"], key=s["finals"].get) if s["finals"] else None
        if final_source:
            finals_set.add(final_source)
        result.append({
            "source_name": src,
            "final_source": final_source,
            "objects_count": len(objs),
            "objects_with_violation": sum(1 for o in objs.values() if o["vv"] > 0),
            "overheat_count": sum(1 for o in objs.values() if o["a75"] > 0),
            "tu_count": s["tu"],
            "volume_total": round(s["volume"], 2),
            "violation_volume": round(s["viol"], 2),
            "violation_pct": round(_cap_pct(s["viol"] / s["volume"] * 100), 2) if s["volume"] else 0.0,
        })
    result.sort(key=lambda x: x["violation_pct"], reverse=True)
    return {"period_id": period_id, "sources": result, "final_sources": sorted(finals_set)}


@router.get("/source-dynamics")
def source_dynamics(
    source_name: str,
    object_type: Optional[str] = None,
    is_dead_end: Optional[str] = None,
    system_type: Optional[str] = None,
    db: Session = Depends(get_db),
):
    """% некачества по неделям для одного источника."""
    periods = db.query(ReportPeriod).order_by(ReportPeriod.period_start).all()
    points = []
    for p in periods:
        q = db.query(TuReportRow).filter(TuReportRow.period_id == p.id)
        if source_name == "— без источника":
            q = q.filter((TuReportRow.source_name == None) | (TuReportRow.source_name == ""))  # noqa: E711
        else:
            q = q.filter(TuReportRow.source_name == source_name)
        q = _apply_filters(q, object_type, is_dead_end, system_type, None, None)
        rows = q.all()
        if not rows:
            continue
        vol = sum(r.volume_total or 0 for r in rows)
        viol = sum(_violation_volume(r) for r in rows)
        points.append({
            "period_id": p.id,
            "period_start": p.period_start,
            "period_end": p.period_end,
            "volume_total": round(vol, 2),
            "violation_pct": round(_cap_pct(viol / vol * 100), 2) if vol else 0.0,
        })
    return {"source_name": source_name, "points": points}


@router.get("/weekly-summary/{period_id}/source-objects")
def week_source_objects(
    period_id: int,
    source_name: str,
    object_type: Optional[str] = None,
    is_dead_end: Optional[str] = None,
    system_type: Optional[str] = None,
    db: Session = Depends(get_db),
):
    """Объекты одного источника за неделю (для раскрытия в разбивке по источникам)."""
    q = db.query(TuReportRow).filter(TuReportRow.period_id == period_id)
    if source_name == "— без источника":
        q = q.filter((TuReportRow.source_name == None) | (TuReportRow.source_name == ""))  # noqa: E711
    else:
        q = q.filter(TuReportRow.source_name == source_name)
    q = _apply_filters(q, object_type, is_dead_end, system_type, None, None)
    rows = q.all()

    by_obj: dict = {}
    for r in rows:
        o = by_obj.setdefault(r.object_id, {
            "object_id": r.object_id, "object_name": r.object_name,
            "tu_id": r.tu_id, "volume": 0.0, "viol": 0.0,
        })
        o["volume"] += r.volume_total or 0
        o["viol"] += _violation_volume(r)
        if (r.volume_total or 0) > 0 and not o["tu_id"]:
            o["tu_id"] = r.tu_id

    objects = [{
        "object_id": o["object_id"],
        "object_name": o["object_name"],
        "tu_id": o["tu_id"],
        "volume_total": round(o["volume"], 2),
        "violation_pct": round(_cap_pct(o["viol"] / o["volume"] * 100), 2) if o["volume"] else 0.0,
    } for o in by_obj.values()]
    objects.sort(key=lambda x: x["violation_pct"], reverse=True)
    return {"source_name": source_name, "objects": objects}


@router.get("/weekly-summary/{period_id}/outage-stats")
def week_outage_stats(period_id: int, db: Session = Depends(get_db), device_db: Session = Depends(get_device_db)):
    """Сколько отключений ГВС (прекращений/ограничений) действовало в течение недели."""
    period = db.query(ReportPeriod).get(period_id)
    if period is None:
        raise HTTPException(404, "Период не найден")
    from datetime import datetime as dt, time as dtime
    start = dt.combine(period.period_start, dtime.min)
    end = dt.combine(period.period_end, dtime.max)

    rows = (
        device_db.query(Outage.impact, Outage.address_norm)
        .filter(Outage.service_gvs == True)  # noqa: E712
        .filter(Outage.start_fact != None)  # noqa: E711
        .filter(Outage.start_fact <= end)
        .filter((Outage.end_fact == None) | (Outage.end_fact >= start))  # noqa: E711
        .all()
    )
    counts: dict = {}
    addrs: dict = {}
    for impact, addr in rows:
        key = impact or "иное"
        counts[key] = counts.get(key, 0) + 1
        addrs.setdefault(key, set()).add(addr)

    breakdown = [
        {"impact": k, "count": v, "objects_count": len(addrs.get(k, set()))}
        for k, v in sorted(counts.items(), key=lambda x: -x[1])
    ]
    return {
        "period_id": period_id,
        "total": sum(counts.values()),
        "breakdown": breakdown,
    }
